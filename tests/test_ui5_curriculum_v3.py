from __future__ import annotations
import ast
import contextlib
import copy
import hashlib
import io
import json
import os
import sys
from pathlib import Path
import tempfile
import tarfile
from types import SimpleNamespace
import unittest
from unittest import mock

from eaglevl.train.ui5_curriculum_profiles import curriculum_phases, profile_env
from eaglevl.train.ui5_token_contract import TOKEN_FIELDS, tokenizer_contract
from scripts.ui5_output_validity import classify_output, audit_evaluation, audit_archive
from scripts import ui5_curriculum_v3 as v3
from scripts import summarize_ui5_curriculum_diagnostics as diagnostics
from tests.test_ui5_curriculum_diagnostics import write_jsonl
from tests import test_ui5_detail_audit_restart as fixtures


class ProfileAndValidityTests(unittest.TestCase):
    def test_actual_anchor_worker_accepts_four_of_four_and_reuses_existing_inferencer(self):
        import argparse
        import time
        from typing import Any
        from PIL import Image
        source = v3.ROOT / "scripts/inference_ui_defect_locany.py"
        names = {"run_anchor_inference", "load_hard_task_rows"}
        functions = [n for n in ast.parse(source.read_text(encoding="utf-8")).body
                     if isinstance(n, ast.FunctionDef) and n.name in names]
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "source.png"
            Image.new("RGB", (20, 20)).save(image)
            inferencer = object()
            calls = []
            def prediction(**kwargs):
                self.assertIs(kwargs["inferencer"], inferencer)
                calls.append(kwargs)
                return "<box>none</box>", SimpleNamespace(status="ok"), [], [], {}, []
            namespace = {"argparse": argparse, "TaskConfig": SimpleNamespace, "Any": Any,
                         "Path": Path, "json": json, "time": time, "os": os, "Image": Image,
                         "_load_python_module": lambda *args: None, "set_sample_seed": lambda seed: None,
                         "stable_sample_seed": lambda *args: 42, "assert_lossless_coverage": lambda *args: None,
                         "predict_with_direct_full_image": prediction, "predict_with_lossless_tiles": prediction,
                         "_score_hard_prediction": lambda *args: {"exact_correct": True, "TP_box": 0, "FP_box": 0, "FN_box": 0, "image_confusion": "TN"},
                         "_atomic_write_jsonl": write_jsonl, "atomic_write_json": v3.evaluation.atomic_write_json,
                         "_paired_baseline_crop_rollouts": mock.Mock(side_effect=AssertionError("anchor must not require 0/4"))}
            exec(compile(ast.Module(body=functions, type_ignores=[]), str(source), "exec"), namespace)
            for task in ("occlusion", "content_missing"):
                rows = root / f"{task}_anchors.jsonl"
                write_jsonl(rows, [{"record_id": task, "sample_id": task, "task": task,
                    "crop_complete4": True, "crop_correct_count": 4, "_resolved_image_path": str(image),
                    "prompt": "Locate defects.", "gt_global": [], "_base_plan_width": 20,
                    "_base_plan_height": 20, "_base_tiles": [[0, 0, 20, 20]]}])
                args = SimpleNamespace(anchor_groups_jsonl=rows, expected_anchor_task_count=1,
                    hard_groups_jsonl=None, rollout_bundle_root=root, rollout_scorer_script=root / "scorer.py",
                    hard_rollout_output_dir=root / task / "rollout4", seed=42,
                    hard_rollout_iou_threshold=.1, evaluation_identity_digest="bound")
                with contextlib.redirect_stdout(io.StringIO()):
                    summary = namespace["run_anchor_inference"](args, inferencer, SimpleNamespace(task_name=task))
                self.assertEqual(summary["group_count"], 1)
                self.assertEqual(summary["exact_correct"], 1)
                self.assertEqual("tiles_override" in calls[-1], task != "content_missing")

    def test_real_workbook_guard_accepts_v3_extra_sheets_and_rejects_stale_values(self):
        from tests import test_ui5_curriculum_artifacts as fixtures_artifacts
        from openpyxl import load_workbook
        source = (v3.ROOT / "shell/run_locany_ui5_crop_rollout4_curriculum_h20x2.sh").read_text(encoding="utf-8")
        block = source.split("evaluation_recorded() {", 1)[1].split("evaluation_seconds()", 1)[0]
        guard = block.split("<<'PY'\n", 1)[1].rsplit("\nPY", 1)[0]
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            checkpoint = fixtures_artifacts.make_checkpoint(root / "model", 0)
            artifact = fixtures_artifacts.ArtifactTests()
            result = artifact.update(root, step=0, metrics=fixtures_artifacts.uniform_metrics(.5, .4),
                                     checkpoint=checkpoint, extra_diagnostics={
                                         "provenance": [{"step": 0, "policy": "boundary_v3", "loss": None}],
                                         "output_validity": [{"step": 0, "image_invalid": 1}]})
            arguments = ["guard", result["checkpoints_json"], result["workbook"], "0"]
            with mock.patch.object(sys, "argv", arguments), self.assertRaises(SystemExit) as exited:
                exec(compile(guard, "actual_launcher_guard", "exec"), {})
            self.assertEqual(exited.exception.code, 0)
            workbook = load_workbook(result["workbook"])
            workbook["output_validity"]["B2"] = 99
            workbook.save(result["workbook"])
            workbook.close()
            with mock.patch.object(sys, "argv", arguments), self.assertRaises(SystemExit) as exited:
                exec(compile(guard, "actual_launcher_guard", "exec"), {})
            self.assertEqual(exited.exception.code, 1)

    def test_evidence_export_is_not_misclassified_as_missing_runtime_results(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "evidence.tar.gz"
            with tarfile.open(path, "w:gz") as archive:
                for step in (0, 200):
                    for task in diagnostics.TASKS:
                        summary = {"totals": {"dataset_images": 2, "parse_error": 1, "inference_error": 0}}
                        raw = {"parse": {"status": "parse_error"}, "raw_answer": "<ref>none</box><|im_end|>"}
                        for name, value in ((f"_worker_summaries/{task}.json", summary), (f"ui_{task}/raw/1.json", raw)):
                            data = json.dumps(value).encode()
                            member = tarfile.TarInfo(f"evaluation/step-{step:06d}/{name}")
                            member.size = len(data)
                            archive.addfile(member, io.BytesIO(data))
            report = audit_archive(path)
            self.assertTrue(all(r["invalid_export_complete"] and r["unexported_images"] == 1 for r in report["rows"]))
            self.assertTrue(all(r["worker_totals"]["inference_error"] == 0 for r in report["rows"]))

    def test_v3_profile_reaches_reporting_and_group_sampler(self):
        from eaglevl.train.ui5_curriculum import UI5CurriculumSchedule
        from eaglevl.train.ui5_curriculum_artifacts import _curriculum_phase
        with mock.patch.dict(os.environ, profile_env("global_replay_v3")):
            self.assertEqual(curriculum_phases()[0], (.20, .20, .60, 1e-6))
            schedule = UI5CurriculumSchedule.from_environment(
                {**profile_env("global_replay_v3"), "CURRICULUM_MODE": "scheduled"}, default_total_steps=1200)
            self.assertEqual(schedule.stages[0].pool_weights, (.20, .20, .60))
            for step, ratios in ((400, (.20, .20, .60)), (401, (.25, .25, .50)), (801, (.30, .30, .40))):
                self.assertEqual(diagnostics._phase_for_step(step, 1200)[1][:3], ratios)
                self.assertEqual(_curriculum_phase(step, total_steps=1200)[1][:3], ratios)
        self.assertEqual(profile_env("global_replay_v3")["HARD_RATIOS"], "0.20,0.25,0.30")

    def test_raw_failure_categories_never_rewrite_predictions(self):
        cases = [(None, "missing_raw"), ("", "empty_answer"), ("<|im_end|>", "empty_answer"),
                 ("<ref>cropping</ref>", "ref_only"), ("<box><1><2>", "box_unclosed"),
                 ("<box>bad</box>", "malformed_box")]
        for raw, expected in cases:
            self.assertEqual(classify_output(raw, "parse_error"), expected)
        self.assertEqual(classify_output("<box>none</box>", "ok"), "valid")
        self.assertEqual(classify_output("", "parse_error", runtime_error="oom"), "runtime_error")
        self.assertEqual(classify_output("<ref>x</ref>", "parse_error", {"termination_reason": "length_limit"}), "length_truncated")
        # Text length alone cannot establish a token-budget stop in historical raw.
        self.assertEqual(classify_output("x" * 4096, "parse_error"), "other_invalid")

    def test_losses_stay_missing_and_actual_sampling_is_reported(self):
        with tempfile.TemporaryDirectory() as folder:
            state = Path(folder) / "trainer_state.json"
            state.write_text(json.dumps({"global_step": 200, "log_history": [{"step": 200,
                "loss": 1.2, "loss_attention": None, "weighted_gate_loss": .12,
                "curriculum_hard_samples": 22, "curriculum_anchor_samples": 19,
                "curriculum_global_replay_samples": 59}]}))
            rows = diagnostics.train_curve_rows(step=200, total_steps=1200, trainer_state=state)
            self.assertIsNone(rows[0]["loss_attention"])
            self.assertIsNone(rows[0]["loss_lm"])
            self.assertEqual(rows[0]["weighted_gate_loss"], .12)
            self.assertEqual(rows[0]["actual_hard_ratio_cumulative"], .22)

    def test_anchor_retention_depends_on_actual_predictions_and_paired_seed(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for task in diagnostics.TASKS:
                baseline = {"record_id": task, "sample_seed": 42, "exact_correct": True,
                            "parse_status": "ok", "runtime_error": None, "latency_seconds": 1.,
                            "TP_box": 1, "FP_box": 0, "FN_box": 0, "image_confusion": "TP"}
                current = {**baseline, "exact_correct": False, "parse_status": "parse_error", "TP_box": 0, "FN_box": 1, "image_confusion": "FN"}
                for step, row in ((0, baseline), (200, current)):
                    write_jsonl(root / f"step-{step:06d}/ui_{task}/anchor_inference/predictions.jsonl", [row])
            rows = diagnostics.anchor_retention_rows(step=200, evaluation_dir=root / "step-000200")
            self.assertTrue(all(r["retention_rate"] == 0 and r["delta"] == -1 for r in rows))
            path = root / "step-000200/ui_occlusion/anchor_inference/predictions.jsonl"
            current["record_id"] = "occlusion"
            current["sample_seed"] = 43
            write_jsonl(path, [current])
            with self.assertRaisesRegex(ValueError, "seeds differ"):
                diagnostics.anchor_retention_rows(step=200, evaluation_dir=root / "step-000200")

    def test_private_model_view_does_not_copy_optimizer_or_modify_original(self):
        with tempfile.TemporaryDirectory() as folder:
            source, target = Path(folder) / "checkpoint-12000", Path(folder) / "initial_model"
            source.mkdir()
            (source / "config.json").write_text('{}')
            (source / "model.safetensors").write_bytes(b"weights")
            (source / "optimizer.pt").write_bytes(b"do-not-restore")
            (source / "training_args.bin").write_bytes(b"do-not-restore")
            (source / "trainer_state.json").write_text('{"global_step":12000}')
            report = v3.model_view(source, target)
            self.assertFalse(report["optimizer_restored"])
            self.assertFalse((target / "optimizer.pt").exists())
            self.assertFalse((target / "trainer_state.json").exists())
            (target / "config.json").write_text('{"patched":true}')
            self.assertEqual((source / "config.json").read_text(), '{}')
            self.assertTrue(os.path.samefile(source / "model.safetensors", target / "model.safetensors"))


class TokenContractTests(unittest.TestCase):
    def test_actual_mtp_negative_supervision_keeps_eos_separate_from_null_padding(self):
        import numpy as np
        import torch
        source = v3.ROOT / "eaglevl/train/locany_finetune_magi_stream.py"
        function = next(n for n in ast.walk(ast.parse(source.read_text(encoding="utf-8")))
                        if isinstance(n, ast.FunctionDef) and n.name == "get_targets_flag_with_mtp")
        ids = {"</box>": 3, "</ref>": 6, "<|im_end|>": 10, "<null>": 9,
               "<text_mask>": 16, "<|im_start|>": 11, "assistant": 14}
        tokenizer = SimpleNamespace(convert_tokens_to_ids=lambda text: ids[text], pad_token_id=0)
        obj = SimpleNamespace(processor=SimpleNamespace(tokenizer=tokenizer), block_size=6, ds_name="negative")
        namespace = {"torch": torch, "np": np, "Dict": dict, "IGNORE_TOKEN_ID": -100}
        exec(compile(ast.Module(body=[function], type_ignores=[]), str(source), "exec"), namespace)
        inputs = torch.tensor([11, 15, 12, 13, 10, 12, 11, 14, 12, 2, 4, 3, 10, 12])
        output = namespace["get_targets_flag_with_mtp"](obj, inputs)
        self.assertEqual(output["labels"][15:21].tolist(), [2, 4, 3, 9, 9, 9])
        self.assertEqual(output["labels"][21:27].tolist(), [10, 9, 9, 9, 9, 9])
        self.assertEqual(output["labels"][9:13].tolist(), [2, 4, 3, 10])
        from eaglevl.train.ui5_token_contract import negative_mtp_contract
        tokenizer.encode = lambda *a, **kw: [2, 4, 3]
        self.assertTrue(negative_mtp_contract(inputs, output, tokenizer, 6)["valid"])
        output["labels"][21] = 9
        with self.assertRaisesRegex(ValueError, "separate EOS"):
            negative_mtp_contract(inputs, output, tokenizer, 6)

    def setup_tokenizer(self):
        tokens = {token: index + 1 for index, token in enumerate(TOKEN_FIELDS.values())}
        tokens.update({"<null>": 9, "<|im_end|>": 10})
        tokens.update({f"<{i}>": 100 + i for i in range(1001)})
        class Tokenizer:
            eos_token_id = 10
            def encode(self, text, add_special_tokens=False):
                if text == "<box>none</box>":
                    return [tokens["<box>"], tokens["none"], tokens["</box>"]]
                return [tokens[text]]
        config = SimpleNamespace(**{field: tokens[token] for field, token in TOKEN_FIELDS.items()},
                                 text_config=SimpleNamespace(null_token_id=9, eos_token_id=10))
        return Tokenizer(), config

    def test_full_negative_eos_null_and_coordinate_contract(self):
        tokenizer, config = self.setup_tokenizer()
        self.assertTrue(tokenizer_contract(tokenizer, config, tokenizer)["valid"])
        config.none_token_id = 4064
        report = tokenizer_contract(tokenizer, config, tokenizer)
        self.assertFalse(report["valid"])
        self.assertIn("none_token_id", " ".join(report["errors"]))

    def test_no_silent_processor_token_id_remapping(self):
        tokenizer, config = self.setup_tokenizer()
        other, _ = self.setup_tokenizer()
        other.encode = lambda *args, **kwargs: [999]
        self.assertFalse(tokenizer_contract(tokenizer, config, other)["valid"])


class DecoderObservationTests(unittest.TestCase):
    def test_evidence_shaped_parallel_ref_is_not_committed_as_a_complete_text_block(self):
        import importlib.util
        import torch
        source = v3.ROOT / "eaglevl/utils/locany/generate_utils.py"
        spec = importlib.util.spec_from_file_location("tested_generation_utils", source)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        ids = {"box_start_token_id": 2, "box_end_token_id": 3, "none_token_id": 4,
               "ref_start_token_id": 5, "ref_end_token_id": 6, "null_token_id": 9,
               "im_end_token_id": 10, "coord_start_token_id": 100, "coord_end_token_id": 1100}
        # Exact token structure from the supplied raw, without using any GT.
        block = torch.tensor([5, 4, 3, 9, 9, 9])
        self.assertEqual(module.handle_pattern(block, ids)["tokens"], [5, 4, 3])
        fixed = module.handle_pattern(block, ids, "hybrid", "boundary_v3")
        self.assertEqual(fixed["tokens"], [5])
        self.assertEqual(fixed["type"], "text_ar")
        self.assertFalse(fixed["is_terminal"])
        null = module.handle_pattern(torch.tensor([9] * 6), ids, "hybrid", "boundary_v3")
        self.assertEqual(null["tokens"], [9])  # Not converted to EOS/empty.
        incomplete_none = module.handle_pattern(torch.tensor([2, 4, 9, 9, 9, 9]), ids, "hybrid", "boundary_v3")
        self.assertEqual(incomplete_none["tokens"], [2, 4])  # No closing token fabricated.
        complete = torch.tensor([2, 101, 102, 103, 104, 3])
        self.assertEqual(module.handle_pattern(complete, ids, "hybrid", "boundary_v3")["tokens"], complete.tolist())

    def test_fixed_hybrid_ar_stops_only_on_eos_and_preserves_unexpected_tokens(self):
        import torch
        source = v3.ROOT / "eaglevl/utils/locany/modeling_locateanything.py"
        function = next(n for n in ast.walk(ast.parse(source.read_text()))
                        if isinstance(n, ast.FunctionDef) and n.name == "_sample_token_in_ar")
        ids = {"box_end_token_id": 3, "coord_start_token_id": 100, "coord_end_token_id": 1100,
               "none_token_id": 4, "im_end_token_id": 10, "ref_end_token_id": 6}
        for token in (0, 3, 4, 6, 9, 10, 20, 100):
            namespace = {"self": SimpleNamespace(token_ids=ids), "generation_mode": "hybrid",
                         "decoder_policy": "boundary_v3", "generate_kwargs": {},
                         "sample_tokens": lambda *args, **kwargs: (None, None, torch.tensor([[token]]), None)}
            exec(compile(ast.Module(body=[function], type_ignores=[]), str(source), "exec"), namespace)
            kind, emitted, decision = namespace["_sample_token_in_ar"](
                torch.tensor([[1]]), SimpleNamespace(logits=torch.zeros((1, 1, 1101))))
            self.assertEqual(kind == "im_end", token == 10)
            self.assertEqual(emitted.tolist(), [token])
            self.assertEqual(kind == "boundary_ar", token in (3, 6))

    def test_hybrid_and_slow_record_actual_ar_decision_without_token_repair(self):
        import torch
        source = v3.ROOT / "eaglevl/utils/locany/modeling_locateanything.py"
        tree = ast.parse(source.read_text(encoding="utf-8"))
        function = next(node for node in ast.walk(tree) if isinstance(node, ast.FunctionDef)
                        and node.name == "_sample_token_in_ar")
        token_ids = {"box_end_token_id": 3, "coord_start_token_id": 100, "coord_end_token_id": 1100,
                     "none_token_id": 4, "im_end_token_id": 10}
        for mode, token, expected_type, reason in (("hybrid", 20, "im_end", "hybrid_unexpected_ar_token"),
                ("slow", 20, "continue_ar", "continue_ar"), ("hybrid", 10, "im_end", "sampled_eos"),
                ("hybrid", 3, "box_end_ar", "box_end_ar")):
            namespace = {"self": SimpleNamespace(token_ids=token_ids), "generation_mode": mode,
                         "decoder_policy": "legacy", "generate_kwargs": {},
                         "sample_tokens": lambda *args, **kwargs: (None, None, torch.tensor([[token]]), None)}
            exec(compile(ast.Module(body=[function], type_ignores=[]), str(source), "exec"), namespace)
            kind, emitted, decision = namespace["_sample_token_in_ar"](
                torch.tensor([[1]]), SimpleNamespace(logits=torch.zeros((1, 1, 1101))))
            self.assertEqual(kind, expected_type)
            self.assertEqual(emitted.tolist(), [token])
            self.assertEqual(decision["sampled_token_ids"], [token])
            self.assertEqual(decision["reason"], reason)


class V3SubmissionTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.DetailAuditRestartTests()
        self.fixture.setUp()
        (self.fixture.fixture.project / "jobs/ui5_crop_curriculum_v3_h20x2.yaml").write_text(
            (v3.ROOT / "jobs/ui5_crop_curriculum_v3_h20x2.yaml").read_text(encoding="utf-8"), encoding="utf-8")
        self.addCleanup(self.fixture.doCleanups)
        self.env = self.fixture.state["runtime"]
        self.output = self.fixture.output
        self.calls = self.fixture.fixture.calls
        self.calls.clear()
        manifest_path = Path(self.env["CURRICULUM_DATA_DIR"]) / "curriculum_manifest.json"
        manifest = json.loads(manifest_path.read_text())
        manifest["matched_anchor_groups"] = manifest["hard_groups"]
        manifest["bundle_group_selection"] = {"selected_pool_strata": {
            pool: {task: {"positive": 1, "negative": 1} for task in v3.evaluation.TASKS}
            for pool in ("hard", "matched_anchor", "global_replay")}}
        from scripts.ui5_curriculum_text_revision import metadata, RECIPE, SIDECARS
        source = manifest_path.parent
        recipe = {}
        manifest["outputs"] = {"recipe": str(source / RECIPE)}
        manifest["pools"] = {}
        for pool in ("hard", "matched_anchor", "global_replay"):
            manifest["pools"][pool] = {"training_records": 1}
            recipe[pool] = {"curriculum_pool": pool, "root": "", "annotation": [pool + ".jsonl"],
                            "paths_relative_to_meta": True}
            write_jsonl(source / (pool + ".jsonl"), [{"_ui5_sample_id": pool, "_ui5_task": "occlusion",
                "_ui5_record_kind": "crop", "_ui5_crop_gt_local_1000": [],
                "messages": [{"role": "assistant", "content": "<ref>overlapping elements</ref><box>none</box>"}]}])
        (source / RECIPE).write_text(json.dumps(recipe))
        for filename in SIDECARS:
            write_jsonl(source / filename, [{"sample_id": "unchanged"}])
        manifest.pop("identity_digest")
        manifest["identity_digest"] = hashlib.sha256(json.dumps(manifest, ensure_ascii=False, sort_keys=True,
                                                              separators=(",", ":")).encode()).hexdigest()
        manifest_path.write_text(json.dumps(manifest))
        manifest_path.with_name("_SUCCESS.json").write_text(json.dumps({"complete": True,
            "identity_digest": manifest["identity_digest"], "recipe_sha256": metadata(source / RECIPE)["sha256"],
            "files": {n: metadata(source / n) for n in (RECIPE, *SIDECARS, "hard.jsonl", "matched_anchor.jsonl", "global_replay.jsonl")}}))
        original = self.fixture.fixture.workspace / v3.ORIGINAL_MODEL_RELATIVE
        (original / "config.json").write_text('{}')
        (original / "model.safetensors").write_bytes(b"source-weights")
        inputs = Path(self.env["EVAL_INPUT_DIR"])
        inputs.mkdir()
        for task, filename in v3.evaluation.TASK_GT_FILE.items():
            image = inputs / (task + ".png")
            image.write_bytes(b"not-decoded-in-preparation")
            write_jsonl(inputs / filename, [{"images": [str(image)], "answer": "untouched GT"}])
            for step in (0, 200):
                root = self.output / f"evaluation/step-{step:06d}"
                v3.evaluation.atomic_write_json(root / f"_worker_summaries/{task}.json",
                    {"totals": {"dataset_images": 1, "parse_error": int(step == 200), "inference_error": 0}})
                v3.evaluation.atomic_write_json(root / f"ui_{task}/raw/{task}.json",
                    {"image_path": str(image), "raw_answer": "<ref>task</ref>" if step else "<box>none</box>",
                     "parse": {"status": "parse_error" if step else "ok"}})
        v3.evaluation.atomic_write_json(self.output / "evaluation/step-000200/evaluation_manifest.json",
                                       {"candidate": {"path": str(self.output / "missing-checkpoint")}})

    def execute(self, submit=True):
        args = SimpleNamespace(previous_submission_dir=self.fixture.state_path.parent,
                               degraded_checkpoint=None, submit=submit, mlx_bin="mlx")
        with mock.patch.object(v3, "ROOT", self.fixture.fixture.project), \
             mock.patch.object(v3.preparation, "PROJECT_ROOT", self.fixture.fixture.project), \
             mock.patch.object(v3, "check_storage", return_value={"user_quota_available_bytes": None}), \
             mock.patch.object(v3.shutil, "which", return_value="/bin/mlx"), \
             mock.patch.object(v3.subprocess, "check_output", return_value="e" * 40), \
             mock.patch.object(v3.subprocess, "run", side_effect=self.fixture.fixture.run_command), \
             contextlib.redirect_stdout(io.StringIO()):
            return v3.prepare(args)

    def test_new_run_reuses_publication_and_submits_fixed_h20_profile(self):
        previous = self.fixture.state_path.read_bytes()
        job_path = self.execute()
        job = v3.preparation.yaml.safe_load(job_path.read_text())
        env = job["jobRunParams"]["envsList"]
        self.assertNotEqual(env["OUTPUT_DIR"], self.env["OUTPUT_DIR"])
        for key in ("FROZEN_SELECTION", "PROCESSOR_PATH"):
            self.assertEqual(env[key], self.env[key])
        self.assertNotEqual(env["CURRICULUM_DATA_DIR"], self.env["CURRICULUM_DATA_DIR"])
        new_recipe = json.loads(Path(env["META_PATH"]).read_text())
        for entry in new_recipe.values():
            annotation = Path(entry["annotation"][0])
            self.assertEqual(annotation.parent, Path(env["CURRICULUM_DATA_DIR"]))
            self.assertEqual(json.loads(annotation.read_text())["messages"][0]["content"], "<box>none</box>")
        self.assertEqual(v3.sha(Path(env["META_PATH"])), env["UI5_TRAIN_RECIPE_SHA256"])
        for key, value in profile_env("global_replay_v3").items():
            self.assertEqual(env[key], value)
        self.assertEqual(env["ATTN_IMPLEMENTATION"], "sdpa")
        self.assertEqual(env["MAX_SEQ_LENGTH"], "7268")
        self.assertNotIn("RESUME_FROM_CHECKPOINT", env)
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(self.calls[0][0][:3], ["/bin/mlx", "job", "submitv2"])
        self.assertEqual(previous, self.fixture.state_path.read_bytes())
        self.assertIn("run_ui5_crop_curriculum_v3_h20x2.sh", job["jobRunParams"]["entrypointFullScript"])
        self.calls.clear()
        with self.assertRaisesRegex(RuntimeError, "already reserved"):
            self.execute()
        self.assertEqual(self.calls, [])

    def test_raw_is_required_before_submission_not_replaced_with_f1_guesses(self):
        (self.output / "evaluation/step-000200/_worker_summaries/occlusion.json").unlink()
        with self.assertRaisesRegex(ValueError, "worker summary missing"):
            self.execute()
        self.assertEqual(self.calls, [])

    def test_paired_comparison_uses_six_formal_cells_and_separate_inference_gains(self):
        job = v3.preparation.yaml.safe_load(self.execute(submit=False).read_text())
        env = job["jobRunParams"]["envsList"]
        env["UI5_V3_DEGRADED_MODEL"] = str(self.output / "explicit-comparison-only")
        commands = []
        def execute(command, **kwargs):
            if "run_ui5_curriculum_evaluation.py" not in command[1]:
                return
            commands.append(command)
            destination = Path(command[command.index("--output-dir") + 1])
            v3.evaluation.atomic_write_json(destination / "ui5_metrics.json", {
                "overall": {"image_macro_f1": .5, "bbox_macro_f1": .4, "joint_score": .45}})
        with mock.patch.dict(os.environ, env), mock.patch.object(v3.subprocess, "run", side_effect=execute), \
             mock.patch.object(v3, "audit_evaluation", return_value={"rows": [{"image_invalid": 1}]}), \
             contextlib.redirect_stdout(io.StringIO()):
            v3.compare()
        self.assertEqual(len(commands), 6)
        for command in commands:
            self.assertEqual(command[command.index("--evaluation-purpose") + 1], "decoder_comparison")
            self.assertEqual(command[command.index("--seed") + 1], "42")
            self.assertNotIn("--anchor-groups-jsonl", command)
            self.assertEqual(command[command.index("--detector-crop-manifest") + 1], env["EVAL_DETECTOR_MANIFEST"])
        report = json.loads((Path(env["OUTPUT_DIR"]) / "diagnostics/decoder_comparison.json").read_text())
        self.assertTrue(report["complete"])
        self.assertEqual(len(report["inference_fix_gain"]), 2)
        self.assertEqual(report["training_gain_reference"], "new full UI5 step-0")

    def test_subset_is_image_identity_selected_and_leaves_all_gt_unchanged(self):
        job = v3.preparation.yaml.safe_load(self.execute(submit=False).read_text())
        selected = Path(job["jobRunParams"]["envsList"]["UI5_V3_COMPARISON_INPUT"])
        for task, filename in v3.evaluation.TASK_GT_FILE.items():
            rows = v3.evaluation.read_jsonl(selected / filename)
            self.assertEqual(rows[0]["answer"], "untouched GT")
        self.assertEqual(self.calls, [])

    def test_corrected_text_statistics_and_identity_reach_real_workbook(self):
        from tests import test_ui5_curriculum_artifacts as artifacts
        from openpyxl import load_workbook
        job = v3.preparation.yaml.safe_load(self.execute(submit=False).read_text())
        env = job["jobRunParams"]["envsList"]
        run_root = Path(env["OUTPUT_DIR"])
        evaluation = run_root / "evaluation/step-000000"
        v3.evaluation.atomic_write_json(evaluation / "evaluation_manifest.json", {
            "curriculum_profile": "global_replay_v3", "generation": {"decoder_policy": "boundary_v3"},
            "frozen_selection": {"summary_sha256": "f" * 64}})
        v3.evaluation.atomic_write_json(evaluation / "ui5_metrics.json", {
            "overall": {"image_macro_f1": .5, "bbox_macro_f1": .4, "joint_score": .45}})
        v3.evaluation.atomic_write_json(run_root / "diagnostics/decoder_comparison.json", {
            "complete": True, "cells": [], "inference_fix_gain": []})
        with mock.patch.dict(os.environ, env), \
             mock.patch.object(diagnostics, "hard_transition_rows", return_value=[]), \
             mock.patch.object(diagnostics, "anchor_retention_rows", return_value=[]), \
             mock.patch.object(diagnostics, "anchor_data_coverage_rows", return_value=[]), \
             mock.patch("scripts.ui5_output_validity.audit_evaluation", return_value={"rows": [{"image_invalid": 1}]}):
            outputs = diagnostics.run(SimpleNamespace(step=0, evaluation_dir=evaluation,
                curriculum_dir=Path(env["CURRICULUM_DATA_DIR"]), trainer_state=None, total_steps=1200, output_dir=None))
        extra = json.loads(Path(outputs["extra_diagnostics"]).read_text())
        self.assertEqual([r["training_delta"] for r in extra["training_vs_baseline"]], [0, 0, 0])
        self.assertEqual([r["ref_negative_before"] for r in extra["supervision_format"]], [1, 1, 1])
        artifact_fixture = artifacts.ArtifactTests()
        result = artifact_fixture.update(run_root, step=0, metrics=artifacts.uniform_metrics(.5, .4),
            checkpoint=Path(env["MODEL_PATH"]), extra_diagnostics=extra)
        workbook = load_workbook(result["workbook"], read_only=True)
        try:
            for name in ("supervision_format", "training_text_identity", "output_validity", "training_vs_baseline"):
                self.assertIn(name, workbook.sheetnames)
            rows = list(workbook["training_text_identity"].values)
            self.assertIn(env["META_PATH"], rows[1])
            self.assertIn(env["UI5_TRAIN_RECIPE_SHA256"], rows[1])
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()

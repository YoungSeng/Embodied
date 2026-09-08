"""CPU contract tests for the formal UI14 path; no detector/model GPU calls."""
from __future__ import annotations
import contextlib
import io
import json
import subprocess
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT))
from PIL import Image
from eaglevl.ui_task_registry import UI_TASKS, UI9_TASKS, configure_task_registry, validate_registry
from eaglevl.train.ui_defect_data import (
    identify_ui_defect_task, build_task_source_balanced_rotating_plan,
    materialize_task_source_balanced_rotating_indices, is_positive_ui_defect)
from ui14_common import paths_for, write_json, write_jsonl, read_json, read_jsonl, image_identity, file_digest, SCAN_NAME
from ui14_annotations import source_boxes, synthetic_boxes, norm1000, crop_boxes, training_record, main_image
import prepare_ui14_sft as prepare
import prepare_ui5_eval_detector_crops as detector
from ui5_lossless_tiling import generate_detector_scan_plan, tile_bbox_to_global


def make_record(task_id, positive=True, source="source", crop="full"):
    task = UI_TASKS[task_id]
    row = dict(source_image="page.png", source_image_id=source, source_record_id=source, split="train",
               source_dataset=task.source_dataset, source_version="fixture", width=375, height=800)
    return training_record(row, task, f"{source}-{crop}.png", [[10, 10, 40, 40]] if positive else [], 375, 800, crop)


def source_fixture(root):
    entries = {}
    for task in UI9_TASKS:
        entries[task.task_key] = {"source_dataset": task.source_dataset, "source_version": "fixture-v2.1",
            "key": task.task_key, "kind": "synthetic" if task.task_id >= 7 else "annotated",
            "train_records": 1, "test_records": 1,
            "bbox": {"mode": "logical-width", "width": 375} if task.task_id >= 7 else {"mode": "auto"},
            "split_origin": "repaired_global_page" if task.task_id >= 7 else "provided"}
        for split in ("train", "test"):
            folder = root / task.task_key
            folder.joinpath("sample_imgs").mkdir(parents=True, exist_ok=True)
            image = folder / "sample_imgs" / f"figma-{split}.png"
            # Deliberate cross-source and train/test duplicates to verify pixel identity.
            Image.new("RGB", (750, 1600), "white").save(image)
            record = {"id": "r1", "images": [str(image)], "split": split,
                      "FigmaKey": "shared-design", "FigmaNodeID": split,
                      "messages": [{"role": "user", "content": "GT: 7 answers. Compare reference image."}]}
            if task.task_id >= 7:
                record["split"] = "test" if split == "train" else "train"  # repair preserves old metadata
                reference = folder / "sample_imgs" / f"reference-{split}.png"
                Image.new("RGB", (100, 200), "blue").save(reference)
                record.update(ScreenShotURL=f"/sample_imgs/{image.name}", RawImgURL=str(reference),
                              BBoxCanvasWidth=375, Objects=[{"IssueName": "wrong source label", "Location": [{"rect1": [10, 40, 100, 80]}]}])
            else:
                record["objects"] = {"bbox": [[20, 80, 200, 160]], "bbox_type": "real"}
            write_jsonl(folder / f"{split}.jsonl", [record])
    write_json(root / "manifest.json", {"format_version": 1, "datasets": list(entries.values()),
        "repair_history": [{"run_id": "fixture-repair-v2.1"}], "require_page_disjoint": True,
        "split_unit": "global_page_for_synthetic", "preparation_errors": []})
    write_json(root / "repair_summary.json", {"run_id": "fixture-repair-v2.1", "check_status": "PASS",
        "input_records": 18, "quarantined_records": 0, "clipped_records": 0, "clipped_boxes": 0,
        "datasets": [{"dataset": t.task_key, "before_train": 1, "before_test": 1,
                      "after_train": 1, "after_test": 1, "quarantined": 0,
                      "moved_split": 2 if t.task_id >= 7 else 0} for t in UI9_TASKS]})


def detector_fixture(root, task, split):
    """Publish actual v5 geometry/ready marker from fixed raw detector results."""
    p = paths_for(root, task.task_key, split)
    args = detector.parse_args(["--stage", "crop", "--output-dir", str(p["cache"]),
        "--parser-root", str(root), "--input-dir", str(p["detector_input"].parent),
        "--task-input-manifest", str(p["detector_inputs"]), "--data-split", split,
        "--cache-scope", "full_train" if split == "train" else "full_test",
        "--expected-unique-images", "1", "--no-skip-figma", "--visualization-samples", "1", "--resume"])
    rows = detector.prepare_manifest(args)
    detected = [{**row, "text_detections": [{"bbox": [20, 80, 200, 160], "score": .9},
                                             {"bbox": [20, 850, 400, 930], "score": .9}],
                 "icon_detections": [{"bbox": [100, 1200, 400, 1350], "score": .8}]} for row in rows]
    write_jsonl(p["cache"] / "detections/merged/detections.jsonl", detected)
    write_json(p["cache"] / "detections/detector_config.json", {"parser_commit": "fixture",
        "text": {"model_dir": None}, "icon": {"model": "model.pt"}})
    for stage in ("text", "icon"):
        folder = p["cache"] / "detections" / stage
        write_json(folder / "stage_summary.json", {"images": 1, "workers": 1, "runtime": {"python": stage}})
        write_jsonl(folder / "shard_00000.jsonl", [{"image_id": r["image_id"]} for r in rows])
        write_json(folder / "shard_00000.done.json", {"stage": stage, "count": 1,
            "image_id_digest": detector.digest_ids(r["image_id"] for r in rows)})
    return detector.build_scan_crops(args)


class UI14DataTests(unittest.TestCase):
    def test_registry_has_stable_ids_families_views_and_legacy_classes(self):
        rows = [t.to_dict() for t in UI_TASKS]
        self.assertEqual([t.task_id for t in UI_TASKS], list(range(14)))
        self.assertEqual([t.task_id for t in UI_TASKS if t.view_policy == "full_image"], [4, 5, 11])
        self.assertEqual([t.class_id for t in UI_TASKS[:5]], [2, 1, 0, 3, 4])
        self.assertEqual(set(t.family_id for t in UI_TASKS), {0, 1, 2, 3})
        validate_registry(rows, 14)
        rows[7]["task_id"] = 1
        with self.assertRaisesRegex(ValueError, "routing drift"): validate_registry(rows, 14)

    def test_same_prompt_sources_use_explicit_task_ids(self):
        for a, b in ((1, 7), (2, 8), (6, 10)):
            ra, rb = make_record(a), make_record(b)
            self.assertEqual(ra["conversations"][0], rb["conversations"][0])
            self.assertEqual(identify_ui_defect_task(ra)[1], a)
            self.assertEqual(identify_ui_defect_task(rb)[1], b)
        with self.assertRaisesRegex(ValueError, "Conflicting"):
            identify_ui_defect_task({**make_record(7), "task_key": "cropping"})

    def test_synthetic_reference_example_scales_y_by_screenshot_width(self):
        raw = {"BBoxCanvasWidth": 375, "Objects": [{"Location": [{"rect1": [183, 449, 269, 479]}]}]}
        box = source_boxes(raw, 941, 2048, synthetic=True)[0]
        self.assertEqual(norm1000(box, 941, 2048), [488, 550, 717, 587])
        crop = [0, 1000, 941, 1500]
        local, contained = crop_boxes([box], crop)
        self.assertEqual(contained, [0])
        for a, b in zip(tile_bbox_to_global(local[0], crop), box): self.assertAlmostEqual(a, b)
        normalized = norm1000(local[0], 941, 500)
        decoded = [v*(941 if i % 2 == 0 else 500)/1000 for i, v in enumerate(normalized)]
        self.assertLess(max(abs(a-b) for a,b in zip(tile_bbox_to_global(decoded,crop), box)), .5)

    def test_rectangle_priority_and_multibox_location_list(self):
        small, big = [1, 2, 3, 4], [5, 6, 8, 9]
        self.assertEqual(synthetic_boxes([{"rect_err": small, "rect_mbr": big}]), [small])
        self.assertEqual(synthetic_boxes({"rect1_shift": small, "rect2_shift": big, "rect_combine": [0,0,10,10]}), [small,big])
        self.assertEqual(synthetic_boxes({"rect2": big, "rect1": small}), [small,big])
        self.assertEqual(source_boxes({"objects": {"bbox": [small,big], "bbox_type": "real"}}, 375,800,synthetic=False), [small,big])
        self.assertEqual(norm1000([10,10,10.001,10.001],750,1600),[13,6,14,7])

    def test_only_explicit_negative_can_use_reference_as_primary(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); image=root/"reference.png"; Image.new("RGB",(10,10)).save(image)
            with self.assertRaises(ValueError): main_image({"RawImgURL":str(image)},root,True)
            with self.assertRaises(ValueError): main_image({"RawImgURL":str(image),"Objects":[]},root,True)
            self.assertEqual(main_image({"LocalImgURL":str(image),"Objects":[]},root,True),image.resolve())

    def test_single_sided_source_rotation_and_mixed_ratio(self):
        for positive in (True,False):
            records = [make_record(6,positive,"many",str(i)) for i in range(7)] + [make_record(6,positive,"one")]
            plan = build_task_source_balanced_rotating_plan(records)
            draws = materialize_task_source_balanced_rotating_indices(plan,seed=42,epoch_index=0)
            self.assertEqual(len(draws),plan["epoch_length"])
            self.assertTrue(all(is_positive_ui_defect(records[i]) == positive for i in draws))
            sources = [records[i]["source_image_id"] for i in draws]
            self.assertLessEqual(abs(sources.count("many")-sources.count("one")),1)
            visited = {i for epoch in range(20) for i in materialize_task_source_balanced_rotating_indices(plan,seed=42,epoch_index=epoch)}
            self.assertEqual(visited,set(range(len(records))))
            self.assertEqual(draws,materialize_task_source_balanced_rotating_indices(plan,seed=42,epoch_index=0))
        records = [make_record(7,True,"p"),make_record(7,False,"n")]
        draws=materialize_task_source_balanced_rotating_indices(build_task_source_balanced_rotating_plan(records),seed=42)
        self.assertEqual(sum(not is_positive_ui_defect(records[i]) for i in draws),2*sum(is_positive_ui_defect(records[i]) for i in draws))

    def test_context_plan_is_input_only_preserves_neighbors_and_raw_edges(self):
        raw = [{"bbox":[10,650,400,730],"source":"text"}, {"bbox":[10,745,400,825],"source":"text"},
               {"bbox":[50,1200,350,1250],"source":"icon"}]
        p = generate_detector_scan_plan(750,1600,raw,task="synth_loneword",target_tile_height=750)
        self.assertIn([650,825],p["protected_vertical_bands"])
        self.assertTrue(all(not 650 < y < 825 for y in p["horizontal_seams"]))
        self.assertTrue(set(p["horizontal_seams"]).issubset(p["safe_raw_detector_edge_candidates"]))
        self.assertFalse(p["gt_used"])
        self.assertEqual(p["processed_pixel_ratio"],1)
        legacy=generate_detector_scan_plan(750,1600,raw,task="text_ellipsis")
        self.assertEqual(legacy["protected_vertical_bands"],[])

    def test_normalization_freezes_splits_and_cache_roundtrip_without_gpu(self):
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            root=Path(tmp); source_fixture(root/"input")
            original = {str(p):file_digest(p) for p in (root/"input").rglob("*.jsonl")}
            args=SimpleNamespace(ui9_data_root=root/"input",output_dir=root/"output",ui5_recipe="audited.json",ui5_test_dir=root/"old_test")
            prepare.normalize(args)
            for path,digest in original.items(): self.assertEqual(file_digest(path),digest)
            self.assertEqual(len(read_json(root/"output/task_registry.json")["tasks"]),14)
            self.assertEqual(len(read_json(root/"output/ui9_image_overlap.json")["train_test"]),1)
            for task in UI9_TASKS:
                for split in ("train","test"):
                    p=paths_for(root/"output",task.task_key,split)
                    rows=list(read_jsonl(p["normalized"]))
                    self.assertEqual(len(rows),1)
                    self.assertEqual(rows[0]["boxes_px"],[[20,80,200,160]])
                    self.assertEqual(rows[0]["split"],split)
                    self.assertEqual(set(next(read_jsonl(p["detector_input"])).keys()),{"image"})
                    if task.view_policy=="crops":
                        detector_fixture(root/"output",task,split)
                        derived=prepare.crop_annotations(root/"output",task,split,rows)
                        self.assertTrue(all(d["split"]==split for d in derived))
                        self.assertTrue(all("GT:" not in d["conversations"][0]["value"] for d in derived))
                        self.assertTrue(all(d["task_id"]==task.task_id for d in derived))
                        prepare.validate_task_cache(root/"output",task,split,1)
            p=paths_for(root/"output","synth_cropping","test")
            with p["detector_input"].open("a") as handle: handle.write("{}\n")
            with self.assertRaises(RuntimeError): prepare.validate_task_cache(root/"output",UI_TASKS[7],"test",1)


class UI14EvaluationTests(unittest.TestCase):
    def test_eight_worker_commands_route_all_14_views_and_keep_new_figma_inputs(self):
        import run_ui5_parallel_inference as parallel
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); manifest=root/"eval.json"
            specs=[{**t.to_dict(),"cache":str(root/t.task_key),"scan_name":SCAN_NAME,"skip_figma":t.task_id<5} for t in UI_TASKS]
            write_json(manifest,{"tasks":specs})
            argv=["parallel","--checkpoint",str(root/"checkpoint-1000"),"--processor-path",str(root/"base"),
                "--input-dir",str(root),"--output-dir",str(root/"pred"),"--gpu-devices","0,1,2,3","--workers-per-gpu","2",
                "--attn-implementation","sdpa","--inference-script",str(ROOT/"scripts/inference_ui_defect_locany.py"),
                "--eval-manifest",str(manifest),"--save-raw-answer"]
            with mock.patch.object(sys,"argv",argv): args=parallel.parse_args()
            self.assertEqual(len(args.tasks),14)
            self.assertEqual(args.workers_per_gpu,2)
            for task in UI_TASKS:
                command=parallel.build_command(args,task.task_key,str(task.task_id%4),root/"summary.json")
                self.assertEqual(command[command.index("--tasks")+1],task.task_key)
                self.assertEqual(command[command.index("--processor-path")+1],str(root/"checkpoint-1000"))
                self.assertEqual("--skip-figma" in command,task.task_id<5)
                self.assertEqual(command[command.index("--inference-crop-mode")+1],"detector_scan" if task.view_policy=="crops" else "full_image")
                self.assertEqual(command.count("--detector-crop-manifest"),int(task.view_policy=="crops"))

    def test_full_evaluation_resume_repairs_missing_ui9_and_keeps_best_ui5(self):
        self._exercise_full_evaluation(1000)

    def test_step_zero_evaluation_writes_all_14_tasks_and_resumes_without_fake_training(self):
        self._exercise_full_evaluation(0)

    def test_real_scorer_retries_preserve_step_zero_reports_and_predictions(self):
        self._exercise_full_evaluation(0, real_scorer=True)

    def test_real_scorer_retries_preserve_step_1000_reports_and_predictions(self):
        self._exercise_full_evaluation(1000, real_scorer=True)

    def _exercise_full_evaluation(self, step, real_scorer=False, eval_set_id=None, external=False):
        import run_ui14_eval as evaluate
        from openpyxl import load_workbook
        from eaglevl.train.ui5_excel_logger import UI5ExcelLogger
        from locany_ui5_common import TASK_JSONL
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()) as console:
            root=Path(tmp); output=root/"run"; checkpoint=output/f"checkpoint-{step}"; specs=[]
            for task in UI_TASKS:
                test=root/(TASK_JSONL[task.task_key] if task.task_id<5 else task.task_key+".jsonl")
                image=str(root/f"page-{task.task_id}.png")
                write_jsonl(test,[dict(image=image,source_image=image,source_image_id=f"id-{task.task_id}",boxes_px=[[10,10,40,40]],objects={"bbox":[[10,10,40,40]]})])
                specs.append({**task.to_dict(),"test":str(test),"split":"test","skip_figma":task.task_id<5,"cache":None,"expected_records":1})
                write_json(output/f"inference-checkpoint-{step}-ui14"/task.task_key/"gate/page.json",dict(
                    image_path=image,prediction_status="defect",final_boxes_pixel_xyxy=[[10,10,40,40]],p_defect=.7,
                    prediction_boxes=1,would_pass=True,coarse_boxes_px=[],coordinate_space="norm1000",image_width=375,image_height=800))
                write_json(output/f"inference-checkpoint-{step}-ui14"/task.task_key/f"page-{task.task_id}_defect.json",
                           [{"bbox_2d": [10,10,40,40], "class_id": task.class_id}])
            if eval_set_id:
                for spec in specs[7:]: spec["parent_positive_test"]=spec["test"]
            manifest=root/"evaluation_manifest.json"; write_json(manifest,{"tasks":specs,"eval_set_id":eval_set_id})
            saved_specs=[{**s,"test":str(root/"old_test"/Path(s["test"]).name)} for s in specs] if external else specs
            write_json(checkpoint/"config.json",dict(ui_num_tasks=14,ui_task_registry=saved_specs))
            recipe=root/"recipe.json"; write_json(recipe,{})
            args=SimpleNamespace(output_dir=output,checkpoint=checkpoint,skip_patch=True,base_model=root/"cpt-9000",
                step=step,project_root=ROOT,eval_gpu_devices="0,1,2,3",attn_implementation="sdpa",scorer_root=ROOT,
                input_dir=root,recipe_path=recipe,tile_nms_iou=.5,external_eval_set=external)
            metric={g:dict(precision=.8,recall=.8,f1=.8,tp=4,fp=1,fn=1,tn=0) for g in ("image","bbox")}
            ui5={"tasks":{t.task_key:metric.copy() for t in UI_TASKS[:5]},"macro":{g:dict(precision=.8,recall=.8,f1=.8) for g in ("image","bbox")}}
            legacy_report = output/f"evaluation/raw/ui14-step-{step}/old-report.txt"
            legacy_report.parent.mkdir(parents=True)
            legacy_report.write_text("Preserve the previous scoring attempt.", encoding="utf-8")
            prediction = output/f"inference-checkpoint-{step}-ui14"
            protected = [legacy_report, checkpoint/"config.json", *prediction.rglob("*.json")]
            snapshots = {p: (p.read_bytes(), p.stat().st_mtime_ns) for p in protected}

            def execute_score(command, *, cwd, stage):
                if stage == "ui14_parallel_inference":
                    return  # Fixtures already contain predictions; no GPU/model calls.
                self.assertEqual(stage, "ui5_score")
                result = subprocess.run(command, cwd=cwd, capture_output=True, text=True,
                                        encoding="utf-8", errors="replace", timeout=60)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

            with mock.patch.dict("os.environ",{"UI_EVAL_MANIFEST":str(manifest),"INIT_CPT_STEP":"9000"}), \
                 mock.patch.object(evaluate,"validate_evaluation_manifest",return_value=specs), \
                 mock.patch("run_ui5_eval.run_checked") as runner, \
                 (contextlib.nullcontext() if real_scorer else
                  mock.patch("collect_ui5_metrics.parse_markdown_report",side_effect=lambda *a: json.loads(json.dumps(ui5)))):
                if step == 0:
                    runner.side_effect = RuntimeError("fixture inference failed")
                    with self.assertRaisesRegex(RuntimeError, "fixture inference failed"):
                        evaluate.run(args)
                    self.assertFalse(evaluate.is_complete(output, step, manifest, checkpoint))
                    self.assertEqual(read_json(output/"evaluation/ui14-step-0.json")["status"], "failed")
                    empty = load_workbook(output/"diagnostics/ui5_training_evaluation.xlsx", read_only=True)
                    try:
                        self.assertEqual([s.max_row for s in empty], [1, 1])
                    finally:
                        empty.close()
                    runner.side_effect = None
                    runner.reset_mock()
                runner.side_effect = execute_score if real_scorer else None
                evaluate.run(args)
                inference_command=runner.call_args_list[0].args[0]
                self.assertEqual(inference_command[inference_command.index("--workers-per-gpu")+1],"2")
                self.assertEqual(inference_command[inference_command.index("--exclusive-gpu-tasks")+1], "synth_loneword")
                self.assertTrue(evaluate.is_complete(output,step,manifest,checkpoint))
                self.assertEqual(read_json(output/"evaluation/best_checkpoints.json")["current_best"]["image"]["image_macro_f1"],1. if real_scorer else .8)
                first_state = read_json(output/f"evaluation/ui14-step-{step}.json")
                first_attempt = Path(first_state["evaluation_run_dir"])
                self.assertTrue((first_attempt/"ui14_metrics.json").is_file())
                first_files = {p: p.read_bytes() for p in first_attempt.rglob("*") if p.is_file()}
                workbook_path = output/"diagnostics/ui5_training_evaluation.xlsx"
                workbook = load_workbook(workbook_path)
                try:
                    self.assertEqual(workbook["train_100steps"].max_row, 1)
                    sheet = workbook["eval_1000steps"]
                    self.assertEqual(sheet.max_row, 37)
                    rows = [dict(zip([c.value for c in sheet[1]], r)) for r in sheet.iter_rows(min_row=2, values_only=True)]
                    self.assertEqual({r["step"] for r in rows}, {step})
                    self.assertEqual({r["sft_step"] for r in rows}, {step})
                    self.assertEqual({r["init_cpt_step"] for r in rows}, {9000})
                    self.assertEqual({r["task"] for r in rows}, {t.diagnostic_name for t in UI_TASKS} | {"five_task_macro", "five_task_micro", "ui9_macro", "ui9_micro"})
                    self.assertIn(f"UI5 full-test scorer metrics: SFT step {step}", console.getvalue())
                    self.assertIn("five_task_macro Image", console.getvalue())
                    self.assertNotIn("synth_cropping image P=", console.getvalue())
                    saved = workbook_path.read_bytes()
                    evaluate.run(args)
                    self.assertEqual(runner.call_count, 2)
                    self.assertEqual(workbook_path.read_bytes(), saved)
                    missing_row = next(i for i, r in enumerate(rows, 2) if r["task"] == "synth_cropping")
                    sheet.delete_rows(missing_row)
                    workbook.save(workbook_path)
                finally:
                    workbook.close()
                self.assertFalse(evaluate.is_complete(output,step,manifest,checkpoint))
                evaluate.run(args)
                self.assertTrue(evaluate.is_complete(output,step,manifest,checkpoint))
                state_path = output/f"evaluation/ui14-step-{step}.json"
                state=read_json(state_path); state["tasks"].pop("synth_cropping")
                write_json(state_path,state)
                self.assertFalse(evaluate.is_complete(output,step,manifest,checkpoint))
                evaluate.run(args)
                self.assertTrue(evaluate.is_complete(output,step,manifest,checkpoint))
                self.assertEqual(runner.call_count,6)
                self.assertEqual(len(read_json(output/"evaluation/evaluation_history.json")),1)
                if real_scorer:
                    saved_workbook = workbook_path.read_bytes()
                    saved_history = (output/"evaluation/evaluation_history.json").read_bytes()
                    state = read_json(state_path); state["tasks"].pop("synth_loneword")
                    write_json(state_path, state)
                    with mock.patch.object(evaluate, "score_ui9", side_effect=RuntimeError("UI9 scoring interrupted")):
                        with self.assertRaisesRegex(RuntimeError, "UI9 scoring interrupted"):
                            evaluate.run(args)
                    failed_state = read_json(state_path)
                    self.assertEqual(failed_state["status"], "failed")
                    self.assertFalse(evaluate.is_complete(output, step, manifest, checkpoint))
                    self.assertTrue((Path(failed_state["evaluation_run_dir"])/"all_tasks_evaluation.txt").is_file())
                    self.assertEqual(workbook_path.read_bytes(), saved_workbook)
                    self.assertEqual((output/"evaluation/evaluation_history.json").read_bytes(), saved_history)
                    evaluate.run(args)
                    self.assertTrue(evaluate.is_complete(output, step, manifest, checkpoint))
                    self.assertEqual(runner.call_count, 10)
                commands = [c.args[0] for c in runner.call_args_list if c.kwargs["stage"] == "ui5_score"]
                attempts = [c[c.index("--run_name")+1] for c in commands]
                self.assertEqual(len(attempts), len(set(attempts)))
                final_state = read_json(state_path)
                self.assertEqual(read_json(output/"evaluation/evaluation_history.json")[0]["evaluation_run_dir"],
                                 final_state["evaluation_run_dir"])
                for path, (content, mtime) in snapshots.items():
                    self.assertEqual(path.read_bytes(), content)
                    self.assertEqual(path.stat().st_mtime_ns, mtime)
                for path, content in first_files.items():
                    self.assertEqual(path.read_bytes(), content)
                if eval_set_id:
                    self.assertEqual(final_state["identity"]["eval_set_id"],eval_set_id)
                    subset=read_json(Path(final_state["evaluation_run_dir"])/"parent_positive_metrics.json")
                    self.assertEqual(len(subset["tasks"]),7)
                    write_json(manifest,{"tasks":specs,"eval_set_id":"different_set"})
                    with self.assertRaisesRegex(ValueError,"Different eval_set_id"):evaluate.run(args)

    def test_finalize_connects_14_streams_original_image_eval_and_bound_report(self):
        from ui5_eval_detector_cache import validate_eval_detector_cache as real_validate
        from locany_ui5_common import TASK_JSONL
        from ui14_profile import validate_prepared_profile
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            root=Path(tmp); data=root/"data"; source_fixture(root/"input")
            audit=root/"ui5/crop_audit_v4_gt_repair"; recipe_path=audit/"training_recipes/old.json"
            legacy_recipe={}; samples=[]
            for task in UI_TASKS[:5]:
                image=root/f"old-{task.task_id}.png"; Image.new("RGB",(375,800),"green").save(image)
                record=make_record(task.task_id,True,f"old-{task.task_id}")
                record["image"]=str(image)
                record["_ui5_crop_source"]="manual_gt_repair" if task.task_id<4 else "full_image"
                record.pop("source_image"); record.pop("_ui5_source_image")
                samples.append(dict(sample_id=record["_ui5_sample_id"],image_id=record["_ui5_image_id"],canonical_path=str(image)))
                train=root/f"oldtrain-{task.task_id}.jsonl"; write_jsonl(train,[record])
                legacy_recipe[task.task_key]={"annotation":str(train),"root":str(root)}
                write_jsonl(root/"old_test"/TASK_JSONL[task.task_key],[{"images":[str(image)],"objects":{"bbox":[]}}])
            write_json(recipe_path,legacy_recipe)
            write_jsonl(audit.parent/"manifest/task_samples.jsonl",samples)
            write_json(root/"ui5_cache"/SCAN_NAME/"eval_detector_cache_ready.json",{"fixture":True})
            write_json(root/"checkpoint-9000/config.json",{"model_type":"locateanything"})
            (root/"checkpoint-9000/model.safetensors").write_bytes(b"fixture only; not model weights")
            args=SimpleNamespace(ui9_data_root=root/"input",output_dir=data,ui5_recipe=recipe_path,
                ui5_test_dir=root/"old_test",ui5_cache=root/"ui5_cache",init_checkpoint=root/"checkpoint-9000")
            prepare.normalize(args)
            for task in UI9_TASKS:
                if task.view_policy=="crops":
                    for split in ("train","test"):
                        detector_fixture(data,task,split)
                        prepare.crop_annotations(data,task,split,list(read_jsonl(paths_for(data,task.task_key,split)["normalized"])))
            def validate(*a,**kw):
                return {} if Path(a[0])==root/"ui5_cache" else real_validate(*a,**kw)
            from ui14_progress import ProgressSession
            with mock.patch("run_ui5_crop_audit.validate_training_ready_marker",return_value={"crop_train_mode":"crop_only"}) as audited, mock.patch("ui5_eval_detector_cache.validate_eval_detector_cache",side_effect=validate), ProgressSession("finalize", data, .1):
                prepare.finalize(args)
                audited.assert_called_once()
                with mock.patch.object(Image, "open", side_effect=AssertionError("repeat finalize opened image")), \
                     mock.patch.object(prepare, "crop_annotations", side_effect=AssertionError("repeat finalize materialized crops")):
                    prepare.finalize(args)
                    validate_prepared_profile({"UI14_DATA_ROOT":str(data)})
            report=read_json(data/"cpu_check_report.json")
            self.assertEqual(read_json(data/"progress/finalize.json")["status"],"completed")
            self.assertFalse(any("progress" in p for p in report["artifact_digests"]))
            self.assertTrue(report["ready"])
            self.assertEqual(report["normalization_resume"]["normalized_records"],18)
            self.assertEqual(set(report["tasks"].values()),{"pass"})
            self.assertEqual(len(report["crop_coverage"]),14)
            self.assertEqual(sum(r["manual_repair_count"] for r in read_json(data/"sampling_stats.json").values()),4)
            self.assertEqual({r["sampling_probability"] for r in read_json(data/"sampling_stats.json").values()},{1/14})
            self.assertTrue(all("evaluation_inputs" in t["test"] and t["skip_figma"] is False for t in read_json(data/"evaluation_manifest.json")["tasks"][5:]))
            self.assertTrue(all("/derived/" in r["annotation"].replace("\\","/") for r in read_json(data/"training_recipe.json").values()))
            validate_prepared_profile({"UI14_DATA_ROOT":str(data)})
            (data/"derived/synth_cropping/train.jsonl").write_text("{}\n",encoding="utf-8")
            with self.assertRaisesRegex(RuntimeError,"artifact changed"): validate_prepared_profile({"UI14_DATA_ROOT":str(data)})

    def test_new_source_scorer_preserves_illegal_output_and_bbox_no_tn(self):
        from run_ui14_eval import score_ui9
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); task=UI_TASKS[7]
            spec={**task.to_dict(),"test":str(root/"test.jsonl")}
            rows=[]
            for index,(boxes,status) in enumerate((([[10,10,40,40]],"defect"),([],"parse_error"),([[10,10,40,40]],"ok"))):
                image=str(root/f"figma-{index}.png")
                rows.append(dict(source_image=image,source_image_id=str(index),boxes_px=boxes))
                write_json(root/task.task_key/"gate"/f"{index}.json",dict(image_path=image,prediction_status=status,
                    final_boxes_pixel_xyxy=boxes if status=="defect" else [],p_defect=.5))
            write_jsonl(spec["test"],rows)
            metrics,_=score_ui9(spec,root,root/"score")
            self.assertEqual({k:metrics["image"][k] for k in ("tp","fp","fn","tn")},{"tp":1,"fp":1,"fn":1,"tn":0})
            self.assertEqual({k:metrics["bbox"][k] for k in ("tp","fp","fn")},{"tp":1,"fp":1,"fn":1})
            self.assertNotIn("tn",metrics["bbox"])
            self.assertEqual(metrics["invalid_pred"],1)
            self.assertEqual(metrics["negative_count"],1)
            write_jsonl(spec["test"],rows[:1])
            self.assertEqual(score_ui9(spec,root,root/"score")[0]["negative_count"],0)

    def test_excel_requires_36_rows_and_keeps_ui5_macro_independent(self):
        from eaglevl.train.ui5_excel_logger import build_eval_rows,UI5ExcelLogger
        metrics={"tasks":{}}
        for t in UI_TASKS:
            value=1.0 if t.task_id<5 else 0.0
            metrics["tasks"][t.task_key]={g:dict(precision=value,recall=value,f1=value,tp=int(value),fp=0,fn=1-int(value),tn=0) for g in ("image","bbox")}
        rows=build_eval_rows(step=1000,checkpoint="checkpoint-1000",metrics=metrics)
        self.assertEqual(len(rows),36)
        self.assertEqual(next(r["f1"] for r in rows if r["task"]=="five_task_macro"),1)
        self.assertEqual(next(r["f1"] for r in rows if r["task"]=="ui9_macro"),0)
        self.assertTrue(all(r["tn"] is None for r in rows if r["granularity"]=="bbox"))
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"ui5_training_evaluation.xlsx"
            legacy=UI5ExcelLogger(path)
            legacy.append_eval(1000,rows[:14])
            logger=UI5ExcelLogger(path,[t.task_key for t in UI_TASKS])
            self.assertFalse(logger.has_eval_step(1000))
            logger.append_eval(1000,rows)
            self.assertTrue(logger.has_eval_step(1000))
            self.assertFalse(logger.append_eval(1000,rows))

    def test_formal_yaml_is_four_a800_and_cpt9000(self):
        from ui14_checks import render_formal_yaml
        with tempfile.TemporaryDirectory() as tmp:
            path,runtime=render_formal_yaml(Path(tmp))
            env=read_json(runtime)
            self.assertIn("compute-3302-yg-cloudnative-ai-aiai.locate-guarantee",path.read_text(encoding="utf-8"))
            self.assertEqual(str(env["INIT_CPT_STEP"]),"9000")
            self.assertEqual(str(env["UI_NUM_TASKS"]),"14")
            self.assertEqual(str(env["GRADIENT_ACCUMULATION_STEPS"]),"2")
            self.assertEqual(str(env["MAX_STEPS"]),"16000")
            self.assertEqual(env["EVAL_FAIL_POLICY"],"stop")
            self.assertEqual(env["EVAL_INFERENCE_WORKERS_PER_GPU"],2)
            self.assertEqual(env["EVAL_AT_START"],1)


if __name__=="__main__": unittest.main()

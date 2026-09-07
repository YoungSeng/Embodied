from copy import deepcopy
import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
import torch
import yaml

from eaglevl.train.ui5_grpo_core import GRPOConfig, digest, file_sha
from eaglevl.train.ui5_grpo_checkpoint import recover, restore_checkpoint, validate_checkpoint
from eaglevl.train.ui5_curriculum_artifacts import normalize_scorer_metrics
from scripts.build_ui5_grpo_mixed_manifest import read_json, write_json
from scripts.run_ui5_grpo_pipeline import verify_evaluation_inputs
from scripts.submit_ui5_grpo_mixed import inherited_source, evaluation_config, render_job, submission_target
from scripts.ui5_grpo_artifacts import detail_tables, write_workbook, register_evaluation
from scripts import run_ui5_curriculum_evaluation as evaluation

ROOT = Path(__file__).resolve().parents[1]


def checkpoint(path, step=100, identity="run"):
    path.mkdir(parents=True)
    files = {"model.safetensors": b"tiny-weights", "config.json": b"{}",
             "rng_sampler_rank0.pt": b"rank0", "rng_sampler_rank1.pt": b"rank1",
             "deepspeed/state/zero_pp_rank_0_mp_rank_00_optim_states.pt": b"opt0",
             "deepspeed/state/zero_pp_rank_1_mp_rank_00_optim_states.pt": b"opt1",
             "deepspeed/state/mp_rank_00_model_states.pt": b"scheduler/model"}
    for name, value in files.items():
        destination = path / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(value)
    write_json(path / "trainer_state.json", dict(global_step=step))
    inventory = {p.relative_to(path).as_posix(): dict(bytes=p.stat().st_size, sha256=file_sha(p))
                 for p in path.rglob("*") if p.is_file()}
    write_json(path / "grpo_complete.json", dict(step=step, run_identity=identity, reference_identity="reference",
               world_size=2, files=inventory, files_digest=digest(inventory)))
    return path


def test_transactional_resume_recovery_and_reject_incomplete_or_corrupt(tmp_path):
    checkpoint(tmp_path / "latest")
    pending = tmp_path / ".pending"
    pending.mkdir()
    (pending / "incomplete").write_text("interrupted ZeRO save")
    assert recover(tmp_path, "run") == 100
    assert not pending.exists()
    checkpoint(pending, 200)
    assert recover(tmp_path, "run") == 200
    assert not (tmp_path / ".previous").exists()
    with pytest.raises(ValueError, match="identity"):
        validate_checkpoint(tmp_path / "latest", "another-reference-run")
    weight = tmp_path / "latest/model.safetensors"
    weight.write_bytes(b"TINY-WEIGHTS")  # same size; hash must still detect it
    with pytest.raises(ValueError, match="corrupted"):
        validate_checkpoint(tmp_path / "latest", "run")


def test_rank_rng_and_group_sampler_restored_with_optimizer_scheduler(tmp_path, monkeypatch):
    import random
    import numpy as np
    latest = tmp_path / "resume/latest"
    latest.mkdir(parents=True)
    random.seed(51)
    np.random.seed(52)
    torch.manual_seed(53)
    state = dict(step=100, run_identity="run", next_draw=401, sampler_identity="pool",
                 sampling_counts={"occlusion/positive/1": 200}, python_rng=random.getstate(),
                 numpy_rng=np.random.get_state(), torch_rng=torch.get_rng_state(), cuda_rng=[])
    torch.save(state, latest / "rng_sampler_rank1.pt")
    write_json(latest / "grpo_complete.json", dict(reference_identity="fixed"))
    expected = random.random(), np.random.random(), torch.rand(3)
    random.random(), np.random.random(), torch.rand(3)
    calls = []
    def load(path, **kwargs):
        calls.append(kwargs)
        return str(latest), dict(run_identity="run", step=100)
    monkeypatch.setattr(torch.cuda, "set_rng_state_all", lambda state: None)
    restored = restore_checkpoint(SimpleNamespace(global_steps=100, load_checkpoint=load),
                                  dict(output_dir=str(tmp_path), identity="run", reference=dict(identity="fixed")), 1)
    assert restored["next_draw"] == 401
    assert calls == [dict(tag="state", load_optimizer_states=True, load_lr_scheduler_states=True)]
    assert random.random() == expected[0]
    assert np.random.random() == expected[1]
    torch.testing.assert_close(torch.rand(3), expected[2])


def test_exact_inherited_v3_submission_and_h20_job_contract(tmp_path):
    previous, v3 = tmp_path / "old", tmp_path / "v3"
    previous.mkdir()
    v3.mkdir()
    env = dict(PROJECT_ROOT="/original", ENV_DIR="/real/conda", PROCESSOR_PATH="/real/processor",
               EVAL_INPUT_DIR="/real/test", EVAL_DETECTOR_MANIFEST="/real/cache", SEED="42")
    old = yaml.safe_load((ROOT / "jobs/ui5_crop_grpo_mixed_v1_h20x2.yaml").read_text())
    old["jobDefVersion"]["imageMeta"]["imageVid"] = "recorded-version"
    old["jobRunParams"]["envsList"] = env
    (v3 / "formal.yaml").write_text(yaml.safe_dump(old))
    source = previous / "snapshot-switch.json"
    write_json(source, dict(snapshot="/frozen/hour021"))
    successor = v3 / "snapshot-switch.json"
    write_json(successor, dict(snapshot="/frozen/hour021", source_submission=str(source),
               runtime=env, job_yaml=str(v3 / "formal.yaml")))
    (previous / "curriculum-v3-text-v3-1.started").write_text(str(successor))
    found, _, job, actual = inherited_source(previous)
    assert found == successor
    run = dict(run_name="mixed-test", project_root="/independent", code_sha="bound-commit",
               output_dir="/new-output", python="/real/conda/bin/python", initial_model="/initial", grpo=GRPOConfig().to_dict())
    rendered = render_job(job, actual, run, "/new-output/run.json", execution_sha="repaired-commit")
    for key in ("resource", "imageMeta", "volumes"):
        assert rendered["jobDefVersion"][key] == old["jobDefVersion"][key]
    runtime = rendered["jobRunParams"]["envsList"]
    assert runtime["ENV_DIR"] == env["ENV_DIR"]
    assert runtime["CODE_REVISION"] == "repaired-commit"
    assert run["code_sha"] == "bound-commit"
    assert all(runtime[key] == "7268" for key in ("MAX_SEQ_LENGTH", "MAX_NUM_TOKENS", "MAX_NUM_TOKENS_PER_SAMPLE"))
    config = evaluation_config(actual)
    assert (config["decoder_policy"], config["generation_mode"]) == ("boundary_v3", "hybrid")
    args = SimpleNamespace(**dict(config, output_dir=tmp_path, checkpoint=Path("/candidate"),
            processor_path=Path("/processor"), input_dir=Path("/test"), python="python",
            worker_script=ROOT / "scripts/inference_ui_defect_locany.py", evaluation_purpose="grpo_full",
            scorer_script=ROOT / "qwen3vl_merge_and_score_fixed_5tasks.py", overwrite=True,
            anchor_groups_jsonl=None, max_images_per_task=0, rollout_bundle_root=Path("/bundle")))
    specs = evaluation.build_worker_specs(args, ("0", "1"), {task: 0 for task in evaluation.TASKS},
                                          Path("/mixed/groups.jsonl"), Path("/identity"))
    assert {s.task: s.physical_gpu for s in specs} == dict(occlusion="0", cropping="0", text_overflow="1", text_ellipsis="1", content_missing="1")
    assert all("--hard-rollout" not in " ".join(s.command) and "--anchor" not in " ".join(s.command) for s in specs)


def test_h20_cluster_switch_changes_only_group_and_queue_and_preserves_run():
    old = yaml.safe_load((ROOT / "jobs/ui5_crop_grpo_mixed_v1_h20x2.yaml").read_text())
    # Inherited runtime resources, including non-template CPU/memory values,
    # must survive the scheduling override and YAML serialization.
    arnold = old["jobDefVersion"]["resource"]["arnoldConfig"]
    arnold["roles"][0].update(cpu=48, memory=500000)
    arnold["groupIds"] = [1234]
    arnold["roles"][0]["queueName"] = "actual-v3-bound-queue"
    env = dict(ENV_DIR="/actual/conda", EVAL_INPUT_DIR="/actual/test")
    run = dict(run_name="same-run", project_root="/independent", code_sha="fixed-run-code",
               output_dir="/same-output", python="/actual/conda/bin/python", initial_model="/fixed-reference",
               grpo=GRPOConfig().to_dict())
    run["identity"] = digest(run)
    original_job, original_run = deepcopy(old), deepcopy(run)
    inherited = render_job(old, env, run, "/same-output/run.json")
    selected = render_job(old, env, run, "/same-output/run.json", cluster="ies_aiai_experience")
    selected = yaml.safe_load(yaml.safe_dump(selected))
    expected = deepcopy(old["jobDefVersion"]["resource"])
    expected["arnoldConfig"]["groupIds"] = [1602]
    expected["arnoldConfig"]["roles"][0]["queueName"] = "compute-329-hl-cloudnative-ai-ies.aiai.experience-guarantee"
    assert selected["jobDefVersion"]["resource"] == expected
    assert inherited["jobDefVersion"]["resource"] == old["jobDefVersion"]["resource"]
    for name in ("imageMeta", "volumes", "gitRepo"):
        assert selected["jobDefVersion"][name] == inherited["jobDefVersion"][name]
    assert selected["namespace"] == inherited["namespace"]
    runtime = selected["jobRunParams"]["envsList"]
    assert runtime["UI5_SUBMISSION_CLUSTER"] == "ies_aiai_experience"
    assert {k: v for k, v in runtime.items() if k != "UI5_SUBMISSION_CLUSTER"} == {
        k: v for k, v in inherited["jobRunParams"]["envsList"].items() if k != "UI5_SUBMISSION_CLUSTER"}
    assert submission_target(selected) == dict(profile="ies_aiai_experience", cluster_id=20,
               group_ids=[1602], queue_name="compute-329-hl-cloudnative-ai-ies.aiai.experience-guarantee")
    assert submission_target(inherited)["group_ids"] == [1234]
    assert old == original_job and run == original_run
    with pytest.raises(ValueError, match="unknown H20"):
        render_job(old, env, run, "/same-output/run.json", cluster="misspelled")


def test_cluster_cli_accepts_resume_code_update_and_rejects_unknown_target(monkeypatch):
    import sys
    from scripts import submit_ui5_grpo_mixed as submitter
    calls = []
    monkeypatch.setattr(submitter, "prepare", lambda args: calls.append(args))
    monkeypatch.setattr(sys, "argv", ["submit_ui5_grpo_mixed.py", "--submit", "--resume-code-update",
                                     "--cluster", "ies_aiai_experience"])
    submitter.main()
    assert calls[-1].submit and calls[-1].resume_code_update
    assert calls[-1].cluster == "ies_aiai_experience"
    monkeypatch.setattr(sys, "argv", ["submit_ui5_grpo_mixed.py", "--submit"])
    submitter.main()
    assert calls[-1].cluster == "default"
    monkeypatch.setattr(sys, "argv", ["submit_ui5_grpo_mixed.py", "--cluster", "misspelled"])
    with pytest.raises(SystemExit) as error:
        submitter.main()
    assert error.value.code == 2
    assert len(calls) == 2


def test_evaluation_input_change_cannot_relabel_step_zero(tmp_path):
    image = tmp_path / "test.png"
    image.write_bytes(b"original")
    inventory = tmp_path / "inputs.json"
    write_json(inventory, {str(image): file_sha(image)})
    run = dict(evaluation_inputs=dict(path=str(inventory), sha256=file_sha(inventory)))
    verify_evaluation_inputs(run)
    image.write_bytes(b"rewritten")
    with pytest.raises(ValueError, match="step 0"):
        verify_evaluation_inputs(run)


def metrics(tp):
    return dict(tasks={task: dict(image=dict(tp=tp, fp=1, fn=1, tn=7),
                      bbox=dict(tp=tp, fp=2, fn=1), invalid_pred=1)
                      for task in evaluation.TASKS})


def test_export_complete_five_task_macro_micro_tables_with_step_zero_delta(tmp_path):
    state = dict(evaluations=[dict(step=step, candidate_checkpoint="verification fixture only",
                metrics=normalize_scorer_metrics(metrics(tp))) for step, tp in ((0, 1), (200, 2))])
    tables = dict(train_curve=[dict(step=10, policy_loss=-.03, kl_loss=.2, replay_loss=1.1)], **detail_tables(state))
    path = tmp_path / "verification_only.xlsx"
    write_workbook(path, tables)
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        assert workbook.sheetnames[-2:] == ["image_detail", "bbox_detail"]
        for name in workbook.sheetnames[-2:]:
            rows = list(workbook[name].values)
            assert len(rows) == 15  # 2 evals x (5 tasks + macro + micro) + header
            headers = rows[0]
            parsed = [dict(zip(headers, row)) for row in rows[1:]]
            assert {row["task"] for row in parsed} == {"ui_" + task for task in evaluation.TASKS} | {"macro", "micro"}
            assert all(row["delta_step0_f1"] == 0 for row in parsed if row["step"] == 0)
            assert all(row["delta_step0_f1"] > 0 for row in parsed if row["step"] == 200)
            if name == "bbox_detail":
                assert all(row["tn"] is None for row in parsed)
    finally:
        workbook.close()


def test_best_aliases_prune_only_superseded_owned_copies_and_repeat_idempotently(tmp_path, monkeypatch):
    from scripts import ui5_grpo_artifacts as writer
    # Windows CI has no symlink privilege. Exercise real checkpoint copy/hash/
    # prune transactions while replacing only this Linux filesystem primitive.
    links = {}
    monkeypatch.setattr(writer, "publish_best_link", lambda alias, target: links.update({alias: target}))
    initial = tmp_path / "initial_model"
    initial.mkdir()
    (initial / "config.json").write_text("{}")
    (initial / "model.safetensors").write_bytes(b"initial fixture")
    mixed = tmp_path / "mixed"
    write_json(mixed / "manifest.json", dict(sampling=[]))
    run = dict(output_dir=str(tmp_path), mixed_dir=str(mixed), identity="run", initial_model=str(initial))
    register_evaluation(run, 0, initial, metrics(1), 1.0)
    assert links[tmp_path / "checkpoints/best-image"] == initial
    candidate200 = checkpoint(tmp_path / "candidate200", 200)
    register_evaluation(run, 200, candidate200, metrics(2), 1.0)
    assert (tmp_path / "checkpoints/step-000200").is_dir()
    candidate400 = checkpoint(tmp_path / "candidate400", 400)
    register_evaluation(run, 400, candidate400, metrics(3), 1.0)
    assert not (tmp_path / "checkpoints/step-000200").exists()
    for name in ("best-image", "best-bbox", "best-joint"):
        assert links[tmp_path / "checkpoints" / name] == tmp_path / "checkpoints/step-000400"
    result = register_evaluation(run, 400, candidate400, metrics(3), 1.0)
    assert result["idempotent"]
    assert read_json(tmp_path / "checkpoints.json")["evaluations"][1]["checkpoint_pruned"]


def test_repaired_runtime_reuses_durable_step_zero_without_gpu_workers(tmp_path, monkeypatch):
    from scripts import run_ui5_grpo_pipeline as pipeline
    from scripts import ui5_grpo_artifacts as writer
    monkeypatch.setattr(writer, "publish_best_link", lambda alias, target: None)
    data = tmp_path / "test"
    data.mkdir()
    image = data / "original.png"
    image.write_bytes(b"unchanged test image fixture")
    for name in evaluation.TASK_GT_FILE.values():
        (data / name).write_text(json.dumps({"image": str(image)}) + "\n")
    detector = data / "detector.jsonl"
    detector.write_text("{}\n")
    inventory = tmp_path / "evaluation_inputs.json"
    write_json(inventory, {str(p): file_sha(p) for p in data.iterdir()})
    initial = tmp_path / "initial_model"
    initial.mkdir()
    weight = initial / "model.safetensors"
    weight.write_bytes(b"unchanged initial checkpoint fixture")
    (initial / "config.json").write_text("{}")
    mixed = tmp_path / "mixed"
    write_json(mixed / "manifest.json", dict(sampling=[]))
    config = dict(input_dir=str(data), detector_crop_manifest=str(detector),
                  decoder_policy="boundary_v3", generation_mode="hybrid")
    run = dict(identity="immutable-run", code_sha="original-code", project_root=str(ROOT),
               output_dir=str(tmp_path), mixed_dir=str(mixed), initial_model=str(initial),
               evaluation=config, reference=dict(identity="fixed-initial-reference"),
               evaluation_inputs=dict(path=str(inventory), sha256=file_sha(inventory)))
    output = tmp_path / "evaluation/step-000000"
    identity = dict(run_identity=run["identity"], step=0, candidate=str(initial),
                    candidate_weights={weight.name: file_sha(weight)},
                    checkpoint_manifest=run["reference"]["identity"], config=config,
                    inputs={name: file_sha(data / name) for name in evaluation.TASK_GT_FILE.values()},
                    detector_manifest=file_sha(detector),
                    scorer=file_sha(ROOT / "qwen3vl_merge_and_score_fixed_5tasks.py"),
                    parser=file_sha(ROOT / "scripts/inference_ui_defect_locany.py"),
                    matcher=file_sha(ROOT / "scripts/ui5_metric_matching.py"))
    write_json(output / "evaluation_manifest.json", identity)
    write_json(output / "ui5_metrics.json", metrics(1))
    write_json(output / "evaluation_status.json", dict(success=True, identity=digest(identity),
               metrics_sha256=file_sha(output / "ui5_metrics.json"), evaluation_seconds=4085.0))
    write_json(tmp_path / "runtime_code_revision.json", dict(execution_code_sha="repaired-code"))
    write_json(tmp_path / "diagnostics/runtime_environment.json",
               dict(ar_numerics="fp32-head-fixed-sdpa-v1", ar_vocab_projection="float32"))
    preserved = {p: p.read_bytes() for p in output.iterdir()}
    def no_worker(*args, **kwargs):
        raise AssertionError("completed step 0 must not rerun a GPU worker")
    monkeypatch.setattr(evaluation, "launch_workers", no_worker)
    monkeypatch.setattr(evaluation, "run_scorer", no_worker)
    pipeline.evaluate(run, 0, initial)
    assert pipeline.evaluate(run, 0, initial)["idempotent"]
    assert {p: p.read_bytes() for p in output.iterdir()} == preserved
    workbook = openpyxl.load_workbook(tmp_path / "diagnostics/ui5_grpo_training_evaluation.xlsx", read_only=True)
    try:
        assert ("runtime_revision.execution_code_sha", "repaired-code") in list(workbook["run_identity"].values)
        assert ("runtime_environment.ar_vocab_projection", "float32") in list(workbook["run_identity"].values)
    finally:
        workbook.close()

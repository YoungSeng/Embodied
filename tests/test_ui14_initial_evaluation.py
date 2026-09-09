"""CPU regressions for diagnostic creation and evaluation-before-training."""
from __future__ import annotations

import contextlib
import io
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from eaglevl.ui_task_registry import UI_TASKS
from eaglevl.train.ui5_excel_logger import (
    EVAL_COLUMNS, EXPECTED_SHEETS, TRAIN_COLUMNS, UI5ExcelLogger, build_eval_rows,
)
from initialize_ui_training_diagnostics import initialize
from tests.test_ui5_excel_logger import training_metrics


class DiagnosticInitializationTests(unittest.TestCase):
    def test_startup_publishes_empty_headers_and_restart_preserves_metrics(self):
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()) as console:
            path = initialize(tmp)
            self.assertIn(str(path), console.getvalue())
            logger = UI5ExcelLogger(path, [t.task_key for t in UI_TASKS])
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual(tuple(workbook.sheetnames), EXPECTED_SHEETS)
                for sheet, columns in zip(workbook, (TRAIN_COLUMNS, EVAL_COLUMNS)):
                    self.assertEqual(tuple(c.value for c in sheet[1]), columns)
                    self.assertEqual(sheet.max_row, 1)
            finally:
                workbook.close()
            self.assertFalse(logger.has_eval_step(0))
            metric = dict(precision=.5, recall=.5, f1=.5, tp=1, fp=1, fn=1)
            metrics = {"tasks": {t.task_key: {g: dict(metric) for g in ("image", "bbox")} for t in UI_TASKS}}
            logger.append_eval(0, build_eval_rows(step=0, checkpoint="checkpoint-0", metrics=metrics))
            train = training_metrics(100)
            train.update(init_cpt_step=9000, init_checkpoint="checkpoint-9000")
            logger.update_train(100, train)
            saved = (path.read_bytes(), path.stat().st_mtime_ns)
            self.assertEqual(initialize(tmp), path)
            self.assertEqual((path.read_bytes(), path.stat().st_mtime_ns), saved)
            self.assertTrue(logger.has_eval_step(0))
            self.assertFalse(logger.update_train(100, train))
            workbook = load_workbook(path, read_only=True)
            try:
                self.assertEqual([s.max_row for s in workbook], [2, 37])
            finally:
                workbook.close()

    def test_initialization_migrates_old_headers_without_losing_records(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "diagnostics.xlsx"
            old = Workbook()
            old.active.title = "train_100steps"
            old.active.append(["step", "epoch", "loss_total"])
            old.active.append([100, .2, 1.5])
            evaluation = old.create_sheet("eval_1000steps")
            evaluation.append(["step", "task", "granularity", "f1", "inference_crop_mode"])
            evaluation.append([0, "text_overflow", "image", .4, "detector_scan"])
            old.save(path)
            old.close()
            logger = UI5ExcelLogger(path)
            self.assertTrue(logger.initialize())
            self.assertFalse(logger.initialize())
            workbook = load_workbook(path, read_only=True)
            try:
                train = dict(zip(TRAIN_COLUMNS, next(workbook["train_100steps"].iter_rows(min_row=2, values_only=True))))
                self.assertEqual((train["step"], train["global_epoch"], train["loss_total"]), (100, .2, 1.5))
                row = dict(zip(EVAL_COLUMNS, next(workbook["eval_1000steps"].iter_rows(min_row=2, values_only=True))))
                self.assertEqual((row["step"], row["f1"], row["eval_inference_crop_mode"]), (0, .4, "detector_scan"))
            finally:
                workbook.close()
            # Partial historical evaluation must still require completion.
            self.assertFalse(logger.has_eval_step(0))


class InitialEvaluationOrderingTests(unittest.TestCase):
    def test_actual_shell_block_evaluates_before_training_reuses_success_and_stops_on_failure(self):
        bash = r"C:\Program Files\Git\bin\bash.exe" if os.name == "nt" else shutil.which("bash")
        if not bash or not Path(bash).is_file():
            self.skipTest("Bash unavailable")
        source = (ROOT / "shell/run_locany_ui5_pipeline.sh").read_text(encoding="utf-8")
        # Execute the real shell orchestration with only GPU/model operations
        # substituted. Completion-set correctness is covered by the UI14 runner.
        functions = source[source.index("run_evaluation() {"):source.index('if [[ "${PIPELINE_MODE}" == "eval" ]]; then')]
        initial = source[source.index('CHECKPOINT_ZERO="'):source.index("while (( current_step < MAX_STEPS )); do")]
        binding = source[source.index("# Publish the UI14 data binding"):source.index('if [[ "${ENABLE_EVAL}" == "0" ]]; then')]
        stubs = r'''
set -e
PIPELINE_PYTHON=fake_python
PROJECT_ROOT=/fixture
OUTPUT_DIR=/fixture/output
BASE_MODEL=/fixture/checkpoint-9000
INIT_CPT_STEP=9000
UI_EVAL_MANIFEST=/fixture/evaluation_manifest.json
UI_TASK_REGISTRY=/fixture/task_registry.json
EVAL_AT_START=1
EVAL_FAIL_POLICY=stop
EVAL_INFERENCE_WORKERS_PER_GPU=2
EVAL_GPU_DEVICES=0,1,2,3
current_step=0
fake_python() {
  case "$1" in
    */export_ui5_checkpoint0.py)
      echo "EXPORTED $*" ;;
    */validate_ui14_ready.py)
      echo "BOUND"
      return "${VALIDATION_EXIT}" ;;
    */ui14_run_recovery.py)
      return "${COMPLETE_STATUS}" ;;
    */run_ui5_eval.py)
      echo "EVALUATED $*"
      return "${INFERENCE_EXIT}" ;;
    *) echo "Unexpected command: $*" >&2; return 99 ;;
  esac
}
'''
        with tempfile.TemporaryDirectory() as tmp:
            script = Path(tmp) / "initial-eval.sh"
            script.write_text(stubs + functions + binding + initial + '\necho "TRAINING_ALLOWED"\n', encoding="utf-8", newline="\n")
            for complete, inference_exit, validation_exit in ((1, 0, 0), (0, 0, 0), (1, 17, 0), (1, 0, 18)):
                with self.subTest(complete=complete, inference_exit=inference_exit, validation_exit=validation_exit):
                    env = {**os.environ, "COMPLETE_STATUS": str(complete), "INFERENCE_EXIT": str(inference_exit), "VALIDATION_EXIT": str(validation_exit)}
                    result = subprocess.run([bash, "--noprofile", "--norc", script.as_posix()], env=env, text=True, capture_output=True)
                    if validation_exit:
                        self.assertEqual(result.returncode, validation_exit)
                        for text in ("EXPORTED", "EVALUATED", "TRAINING_ALLOWED"):
                            self.assertNotIn(text, result.stdout)
                        continue
                    self.assertEqual(result.returncode, inference_exit)
                    self.assertIn("BOUND", result.stdout)
                    if complete:
                        self.assertIn("--init-cpt-step 9000", result.stdout)
                        self.assertIn("--step 0", result.stdout)
                        self.assertIn("--eval-inference-workers-per-gpu 2", result.stdout)
                        self.assertLess(result.stdout.index("BOUND"), result.stdout.index("EXPORTED"))
                        self.assertLess(result.stdout.index("EXPORTED"), result.stdout.index("EVALUATED"))
                        if not inference_exit:
                            self.assertLess(result.stdout.index("EVALUATED"), result.stdout.index("TRAINING_ALLOWED"))
                    else:
                        self.assertNotIn("EXPORTED", result.stdout)
                        self.assertNotIn("EVALUATED", result.stdout)
                    self.assertEqual("TRAINING_ALLOWED" in result.stdout, not inference_exit)
                    if inference_exit:
                        self.assertIn("evaluation failed: step=0", result.stderr)
                        self.assertNotIn("training will continue", result.stderr)


if __name__ == "__main__":
    unittest.main()
